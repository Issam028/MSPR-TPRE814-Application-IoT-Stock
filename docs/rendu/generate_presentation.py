from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "rendu" / "MSPR_TPRE814_Thibault AUTEXIER - Issam HARNOUFI - Zaid ABABOU - Ali WARI_presentation.pptx"
NOTES = ROOT / "docs" / "rendu" / "MSPR_TPRE814_presentation_notes.md"

ASSETS = {
    "architecture": ROOT / "docs" / "rendu" / "presentation_images" / "01_flux_donnees_futurekawa.png",
    "dashboard": ROOT / "docs" / "rendu" / "presentation_images" / "02_dashboard_bresil.png",
    "exploitations": ROOT / "docs" / "rendu" / "presentation_images" / "03_page_exploitations_bresil.png",
    "entrepots": ROOT / "docs" / "rendu" / "presentation_images" / "04_page_entrepots_bresil.png",
    "mcd": ROOT / "docs" / "rendu" / "presentation_images" / "05_mcd_erd_futurekawa.png",
    "wiring": ROOT / "docs" / "rendu" / "presentation_images" / "06_cablage_esp32_dht11.png",
    "thonny": ROOT / "docs" / "rendu" / "presentation_images" / "07_thonny_micropython.png",
    "mqtt": ROOT / "docs" / "rendu" / "presentation_images" / "08_mqtt_bridge_log.png",
    "mail": ROOT / "docs" / "rendu" / "presentation_images" / "09_alerte_mail.png",
    "erp_stock": ROOT / "docs" / "rendu" / "presentation_images" / "10_erp_stock_movements.png",
    "jenkins": ROOT / "docs" / "rendu" / "presentation_images" / "12_jenkins_pipeline_vert.png",
    "docker": ROOT / "docs" / "rendu" / "presentation_images" / "13_docker_services.png",
}

WIDE = (13.333333, 7.5)
NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(30, 64, 175)
SKY = RGBColor(224, 242, 254)
LINE = RGBColor(203, 213, 225)
MUTED = RGBColor(71, 85, 105)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(217, 119, 6)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 252)


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_run(run, size=18, color=NAVY, bold=False, font="Aptos"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_text(slide, text, x, y, w, h, size=18, color=NAVY, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run()
    r.text = text
    set_run(r, size=size, color=color, bold=bold)
    return box


def add_multiline(slide, lines, x, y, w, h, size=16, color=NAVY, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        if bullet:
            p.text = f"• {line}"
        for run in p.runs:
            set_run(run, size=size, color=color)
        p.space_after = Pt(6)
    return box


def add_title(slide, title, kicker=None):
    if kicker:
        add_text(slide, kicker.upper(), 0.62, 0.38, 5.4, 0.28, size=10, color=BLUE, bold=True)
    add_text(slide, title, 0.62, 0.74, 8.8, 0.55, size=28, color=NAVY, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.36), Inches(1.25), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()


def add_footer(slide, number):
    add_text(slide, "FutureKawa - MSPR TPRE814", 0.62, 7.08, 3.2, 0.2, size=8, color=MUTED)
    add_text(slide, str(number), 12.25, 7.08, 0.45, 0.2, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def image_size(path: Path):
    with Image.open(path) as img:
        return img.size


def add_image_fit(slide, path: Path, x, y, w, h, fit="contain"):
    if not path.exists():
        card(slide, x, y, w, h, fill=rgb("fee2e2"), line=rgb("fecaca"))
        add_text(slide, f"Image manquante\n{path.name}", x + 0.2, y + 0.2, w - 0.4, h - 0.4, size=14, color=rgb("991b1b"))
        return None

    iw, ih = image_size(path)
    box_ratio = w / h
    img_ratio = iw / ih
    if fit == "cover":
        if img_ratio > box_ratio:
            height = h
            width = h * img_ratio
        else:
            width = w
            height = w / img_ratio
    else:
        if img_ratio > box_ratio:
            width = w
            height = w / img_ratio
        else:
            height = h
            width = h * img_ratio
    left = x + (w - width) / 2
    top = y + (h - height) / 2
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    return pic


def screenshot_frame(slide, path: Path, x, y, w, h, caption):
    card(slide, x, y, w, h, fill=WHITE, line=LINE)
    add_image_fit(slide, path, x + 0.12, y + 0.12, w - 0.24, h - 0.55, fit="contain")
    add_text(slide, caption, x + 0.2, y + h - 0.36, w - 0.4, 0.22, size=8.8, color=MUTED)


def pill(slide, text, x, y, w, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.36))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    add_text(slide, text, x + 0.1, y + 0.08, w - 0.2, 0.16, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_stat(slide, label, value, x, y, color=BLUE):
    card(slide, x, y, 2.25, 1.15, fill=WHITE, line=LINE)
    add_text(slide, value, x + 0.18, y + 0.18, 1.9, 0.42, size=28, color=color, bold=True)
    add_text(slide, label, x + 0.2, y + 0.7, 1.8, 0.26, size=9.5, color=MUTED)


def flow_node(slide, text, x, y, w=1.75):
    card(slide, x, y, w, 0.7, fill=WHITE, line=rgb("bfdbfe"))
    add_text(slide, text, x + 0.12, y + 0.23, w - 0.24, 0.18, size=11, color=NAVY, bold=True, align=PP_ALIGN.CENTER)


def connector(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = BLUE
    line.line.width = Pt(2)


def notes_text():
    return """# Notes orales - MSPR TPRE814 FutureKawa

Durée cible : 20 minutes. Parle naturellement : l'objectif est de raconter le projet, pas de lire les slides.

1. Introduction : rappeler le besoin FutureKawa et annoncer le fil rouge : stock, IoT, alertes, preuves.
2. Contexte : insister sur la traçabilité et la qualité du stockage du café vert.
3. Solution livrée : expliquer les 3 blocs, siège, pays, IoT.
4. Architecture : décrire le chemin de la donnée de l'entrepôt jusqu'au dashboard.
5. Docker : montrer que tout est reproductible localement.
6. Base de données : expliquer les relations et le dataset réellement trouvé.
7. IoT : dire que le prototype ESP32/DHT11 publie via MQTT.
8. MQTT : expliquer le bridge vers l'API pays.
9. Interface : montrer les pages comme preuve métier.
10. Alertes : expliquer pourquoi une alerte renvoie au contexte d'exploitation.
11. ERP : être honnête : adaptateur simulé POC, pas SAP réel.
12. Tests/Jenkins : insister sur tests automatisés + pipeline vert.
13. Limites : dire ce qui reste industriel : JWT, ERP réel, E2E navigateur, supervision.
14. Conclusion : résumer en une phrase : la chaîne de supervision est démontrable, documentée et vérifiée.
"""


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(WIDE[0])
    prs.slide_height = Inches(WIDE[1])
    blank = prs.slide_layouts[6]

    def new_slide(n, title=None, kicker=None):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        if title:
            add_title(s, title, kicker)
        add_footer(s, n)
        return s

    # 1 cover
    s = new_slide(1)
    add_text(s, "MSPR TPRE814", 0.75, 0.62, 3.2, 0.28, size=12, color=BLUE, bold=True)
    add_text(s, "FutureKawa", 0.75, 1.18, 6.2, 0.7, size=44, color=NAVY, bold=True)
    add_text(s, "Application IoT de supervision des stocks et des conditions de stockage", 0.78, 2.0, 6.3, 0.62, size=20, color=MUTED)
    card(s, 0.78, 3.05, 5.4, 1.78, fill=WHITE, line=rgb("bfdbfe"))
    add_multiline(s, ["Thibault AUTEXIER", "Issam HARNOUFI", "Zaid ABABOU", "Ali WARI"], 1.05, 3.35, 4.5, 1.1, size=16, color=NAVY)
    add_text(s, "Soutenance - 20 minutes", 0.78, 5.12, 3.0, 0.25, size=12, color=BLUE, bold=True)
    screenshot_frame(s, ASSETS["dashboard"], 7.15, 0.75, 5.45, 5.55, "Interface centrale FutureKawa")

    # 2 problem
    s = new_slide(2, "Le problème métier", "Contexte")
    add_multiline(s, [
        "FutureKawa doit suivre ses stocks de café vert dans plusieurs pays.",
        "Les conditions de stockage influencent directement la qualité : température, humidité, ancienneté des lots.",
        "Le siège a besoin d'une vision centralisée, mais les données restent locales par pays.",
    ], 0.72, 1.65, 5.55, 2.1, size=18)
    add_stat(s, "pays couverts", "3", 0.82, 4.25, BLUE)
    add_stat(s, "tables métier", "4", 3.25, 4.25, GREEN)
    add_stat(s, "preuve IoT", "ESP32", 5.68, 4.25, AMBER)
    screenshot_frame(s, ASSETS["dashboard"], 7.35, 1.32, 5.05, 4.55, "Vue siège : indicateurs par pays")

    # 3 solution
    s = new_slide(3, "Ce que nous avons livré", "Périmètre")
    items = [
        ("Frontend React", "dashboard, exploitations, entrepôts, alertes"),
        ("APIs NestJS", "API centrale + API pays réutilisable"),
        ("SQL + Docker", "MySQL par pays, stack reproductible"),
        ("IoT + MQTT", "ESP32/DHT11, Mosquitto, bridge HTTP"),
        ("CI + preuves", "Jenkins, tests, rapport, captures"),
    ]
    y = 1.55
    for title, desc in items:
        card(s, 0.78, y, 5.65, 0.78, fill=WHITE, line=LINE)
        add_text(s, title, 1.02, y + 0.14, 1.9, 0.22, size=15, color=BLUE, bold=True)
        add_text(s, desc, 2.95, y + 0.17, 3.1, 0.2, size=12.5, color=NAVY)
        y += 0.9
    screenshot_frame(s, ASSETS["architecture"], 7.05, 1.25, 5.35, 4.9, "Flux global livré")

    # 4 architecture
    s = new_slide(4, "Architecture applicative", "Flux")
    nodes = [("ESP32", 0.8), ("MQTT", 2.55), ("API pays", 4.3), ("MySQL", 6.05), ("API centrale", 7.8), ("React", 9.85)]
    for label, x in nodes:
        flow_node(s, label, x, 2.0)
    for i in range(len(nodes) - 1):
        connector(s, nodes[i][1] + 1.75, 2.35, nodes[i + 1][1], 2.35)
    add_multiline(s, [
        "La donnée part du capteur, passe par MQTT, puis est persistée dans la base du pays.",
        "L'API centrale interroge les APIs pays pour donner une vue siège.",
        "Chaque pays conserve sa base : Brésil, Colombie, Équateur.",
    ], 0.9, 3.42, 5.4, 1.9, size=17)
    screenshot_frame(s, ASSETS["architecture"], 6.9, 3.15, 5.45, 2.65, "Schéma de flux documenté")

    # 5 docker
    s = new_slide(5, "Déploiement local reproductible", "Docker")
    screenshot_frame(s, ASSETS["docker"], 0.8, 1.45, 6.0, 4.75, "Services Docker et état de la stack")
    add_multiline(s, [
        "3 bases MySQL, une par pays.",
        "3 APIs pays avec la même base de code.",
        "API centrale + frontend web.",
        "Profil dev : Mosquitto, bridge MQTT et simulateur.",
    ], 7.35, 1.65, 4.7, 2.4, size=17, bullet=True)
    add_text(s, "Pourquoi c'est important", 7.35, 4.4, 3.2, 0.26, size=18, color=BLUE, bold=True)
    add_text(s, "Le jury peut relancer le projet sans dépendre d'un environnement caché.", 7.35, 4.85, 4.55, 0.7, size=16, color=NAVY)

    # 6 database
    s = new_slide(6, "Base SQL et jeu de données", "Traçabilité")
    screenshot_frame(s, ASSETS["mcd"], 0.75, 1.42, 6.0, 4.65, "MCD / ERD réel du projet")
    add_stat(s, "lignes lues", "129", 7.15, 1.55, BLUE)
    add_stat(s, "importables", "114", 9.65, 1.55, GREEN)
    add_stat(s, "mesures", "105", 7.15, 2.95, AMBER)
    add_stat(s, "rejetées", "15", 9.65, 2.95, rgb("dc2626"))
    add_text(s, "Point assumé", 7.15, 4.65, 2.5, 0.25, size=17, color=BLUE, bold=True)
    add_text(s, "Le dépôt contient un jeu SQL exploitable dans docs/data_tests. Aucun CSV/Excel officiel n'était présent dans le dépôt. Les 15 rejets correspondent à des mesures sur un entrepôt absent.", 7.15, 5.05, 4.8, 0.82, size=14.5, color=NAVY)

    # 7 IoT
    s = new_slide(7, "Prototype IoT ESP32 + DHT11", "Embarqué")
    screenshot_frame(s, ASSETS["wiring"], 0.8, 1.45, 5.45, 4.5, "Câblage exact : VCC, GND, DATA GPIO32")
    screenshot_frame(s, ASSETS["thonny"], 6.85, 1.45, 5.55, 4.5, "Script MicroPython exécuté dans Thonny")
    add_text(s, "Le capteur publie un JSON température/humidité sur MQTT.", 1.0, 6.25, 10.5, 0.24, size=15, color=NAVY)

    # 8 MQTT
    s = new_slide(8, "MQTT vers API : le pont terrain", "IoT")
    screenshot_frame(s, ASSETS["mqtt"], 0.85, 1.55, 5.85, 4.35, "Bridge MQTT : message reçu puis posté vers l'API")
    add_multiline(s, [
        "Topic : futurekawa/mesures",
        "Payload : id_entrepot, temperature, humidite",
        "Le bridge transforme le message MQTT en POST /mesures.",
        "Le même flux est testable sans matériel avec mosquitto_pub.",
    ], 7.25, 1.65, 4.8, 2.5, size=17, bullet=True)
    pill(s, "Capteur réel ou simulation", 7.25, 4.65, 2.7, BLUE)
    pill(s, "Même contrat JSON", 10.05, 4.65, 2.15, GREEN)

    # 9 UI
    s = new_slide(9, "Interface de supervision", "Frontend")
    screenshot_frame(s, ASSETS["exploitations"], 0.75, 1.36, 5.78, 4.85, "Page exploitations : mesures et graphe lisible")
    screenshot_frame(s, ASSETS["entrepots"], 6.82, 1.36, 5.78, 4.85, "Page entrepôts : détail, lots et historique")

    # 10 alerts
    s = new_slide(10, "Alertes : ne pas afficher un signal isolé", "Décision produit")
    add_multiline(s, [
        "Une alerte doit ramener vers le contexte métier : pays, exploitation, entrepôt, historique.",
        "Le responsable peut comprendre la cause au lieu de voir seulement un badge rouge.",
        "SMTP réel possible ; sinon le mode log garde une preuve démontrable.",
    ], 0.85, 1.55, 5.35, 2.65, size=18)
    screenshot_frame(s, ASSETS["mail"], 6.85, 1.35, 5.45, 4.25, "Preuve notification mail / alerte")
    add_stat(s, "température", "seuil", 0.95, 4.8, AMBER)
    add_stat(s, "humidité", "seuil", 3.4, 4.8, BLUE)

    # 11 ERP/security
    s = new_slide(11, "Adaptateur ERP simulé", "Intégration")
    screenshot_frame(s, ASSETS["erp_stock"], 0.75, 1.42, 6.0, 4.6, "Export stock en format ERP JSON")
    add_multiline(s, [
        "Routes /erp/health, /erp/stock-movements, /erp/quality-alerts.",
        "Protection par clé API et rôle simple.",
        "POC honnête : pas de SAP/Dynamics réel, mais mapping stock/qualité prêt.",
    ], 7.25, 1.65, 4.65, 2.65, size=17, bullet=True)
    pill(s, "stock", 7.25, 4.55, 1.25, BLUE)
    pill(s, "quality", 8.7, 4.55, 1.45, AMBER)
    pill(s, "admin", 10.35, 4.55, 1.25, GREEN)

    # 12 tests
    s = new_slide(12, "Tests et intégration continue", "Validation")
    screenshot_frame(s, ASSETS["jenkins"], 0.75, 1.36, 6.15, 4.7, "Pipeline Jenkins vert")
    add_multiline(s, [
        "Tests automatisés : seuils mesure, mapping ERP, parsing dataset.",
        "Build API pays, API centrale, frontend et image bridge.",
        "Docker Compose validé avant build.",
        "Anomalies et re-tests documentés.",
    ], 7.35, 1.65, 4.55, 2.55, size=17, bullet=True)
    add_text(s, "Message à porter", 7.35, 4.6, 3.2, 0.25, size=18, color=BLUE, bold=True)
    add_text(s, "On ne montre pas seulement une application : on montre une démarche vérifiable.", 7.35, 5.02, 4.45, 0.52, size=16, color=NAVY)

    # 13 limits
    s = new_slide(13, "Limites assumées et suite logique", "Industrialisation")
    cols = [
        ("Sécurité", ["JWT/SSO", "HTTPS", "rate limiting"]),
        ("ERP", ["connecteur réel", "flux retour", "format éditeur"]),
        ("Tests", ["E2E navigateur", "couverture", "SonarQube"]),
        ("Exploitation", ["supervision", "sauvegardes", "déploiement cloud"]),
    ]
    x = 0.78
    for title, lines in cols:
        card(s, x, 1.72, 2.75, 3.7, fill=WHITE, line=LINE)
        add_text(s, title, x + 0.2, 2.0, 2.0, 0.26, size=18, color=BLUE, bold=True)
        add_multiline(s, lines, x + 0.24, 2.55, 2.2, 1.7, size=15, color=NAVY, bullet=True)
        x += 3.08
    add_text(s, "Ces limites ne cachent pas le travail : elles montrent la trajectoire pour passer du POC à une version industrielle.", 1.0, 6.0, 11.2, 0.38, size=16, color=NAVY, align=PP_ALIGN.CENTER)

    # 14 close
    s = new_slide(14, "Conclusion", "À retenir")
    add_text(s, "FutureKawa est une chaîne complète de supervision : donnée terrain, stockage SQL, API, interface, alertes, tests et documentation.", 1.0, 1.55, 10.9, 1.0, size=26, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_stat(s, "pays", "3", 1.55, 3.35, BLUE)
    add_stat(s, "flux IoT", "MQTT", 4.05, 3.35, GREEN)
    add_stat(s, "CI", "Jenkins", 6.55, 3.35, AMBER)
    add_stat(s, "rapport", "59p", 9.05, 3.35, BLUE)
    add_text(s, "Merci", 5.1, 5.65, 3.0, 0.5, size=34, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    return prs


def main():
    prs = build_deck()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    NOTES.write_text(notes_text(), encoding="utf-8")
    print(OUT)
    print(NOTES)


if __name__ == "__main__":
    main()
